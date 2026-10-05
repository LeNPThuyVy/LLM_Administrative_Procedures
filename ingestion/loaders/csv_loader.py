import csv
from pathlib import Path

from openpyxl import load_workbook

from .base_loader import BaseLoader, RawDocument


class CSVLoader(BaseLoader):
    """
    Load dữ liệu từ CSV hoặc XLSX thành các RawDocument có metadata dạng dict.
    """

    def load(self, source_config: dict) -> list[RawDocument]:
        path_value = source_config.get("path")

        if not path_value:
            raise ValueError("CSV source config must contain 'path'")

        path = Path(path_value)

        if not path.exists():
            raise FileNotFoundError(f"Source not found: {path}")

        suffix = path.suffix.lower()

        if suffix == ".csv":
            return self._load_csv(path)

        if suffix == ".xlsx":
            return self._load_xlsx(path)

        raise ValueError(
            f"Unsupported tabular file type: {suffix}"
        )

    def _load_csv(self, path: Path) -> list[RawDocument]:
        documents = []

        with path.open(
            "r",
            encoding="utf-8-sig",
            newline="",
        ) as f:
            reader = csv.DictReader(f)

            if reader.fieldnames is None:
                return []

            reader.fieldnames = [
                str(name).strip()
                for name in reader.fieldnames
            ]

            for row in reader:
                cleaned = {
                    str(key).strip(): (
                        value.strip()
                        if isinstance(value, str)
                        else value
                    )
                    for key, value in row.items()
                    if key is not None
                }

                documents.append(
                    RawDocument(
                        source_path=str(path),
                        raw_text="",
                        metadata=cleaned,
                    )
                )

        return documents

    def _load_xlsx(self, path: Path) -> list[RawDocument]:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)

        try:
            header_row = next(rows)
        except StopIteration:
            workbook.close()
            return []

        headers = [
            str(value).strip()
            if value is not None
            else ""
            for value in header_row
        ]

        documents = []

        for row in rows:
            if all(value is None for value in row):
                continue

            metadata = {}

            for index, value in enumerate(row):
                if index >= len(headers):
                    continue

                key = headers[index]

                if not key:
                    continue

                metadata[key] = (
                    value.strip()
                    if isinstance(value, str)
                    else value
                )

            documents.append(
                RawDocument(
                    source_path=str(path),
                    raw_text="",
                    metadata=metadata,
                )
            )

        workbook.close()
        return documents